#!/usr/bin/env python3
"""One-off import of the old Excel planning workbook into our YAML data files.

Usage:  python import_excel.py PATH/TO/prev.xlsm

Reads the `Recepies` sheet (recipes x ingredients matrix) and the price column of `Ingredient ALL`.
Writes data/stores.yaml, data/ingredients.yaml and data/recipes/<id>.yaml, plus REPORT.md next to this script.

This script is allowed to be messy and is meant to be thrown away once the data has been cleaned by hand.
It only READS the workbook. It never runs its macros and never calculates anything with Excel.
Every guess it makes is listed in REPORT.md.
"""
import re
import sys
import unicodedata
import json
import pathlib
import collections
from decimal import Decimal, ROUND_HALF_UP

import openpyxl

HERE = pathlib.Path(__file__).resolve().parent
ROOT = HERE.parent.parent
DATA = ROOT / "data"

# --- what the Excel calls things ---------------------------------------------------------------
STORES = {  # Excel code -> (our id, display name)
    "SM": ("supermarket", "Supermarket"),
    "AS": ("asian-market", "Asian market"),
    "ONL": ("online", "Online"),
    "DM": ("drugstore", "Drugstore"),
    "NW": ("none", "Not bought"),
    "SM2": ("supermarket", None),  # unknown code, only used by one ingredient; treated as supermarket
}
UNITS = {  # Excel unit -> (our unit, factor from the Excel's kg/L to our g/ml)
    "KG": ("g", 1000),
    "L": ("ml", 1000),
    "PCS": ("pcs", 1),
    "PCK": ("pack", 1),
    "Jar": ("jar", 1),
}
CATEGORIES = {
    "main": "main",
    "brunch": "brunch",
    "helpersnacks": "helper-snack",
    "drink": "drink",
    "misc": "misc",
}
# Names the old sheet itself used before they were abbreviated ("Garlic [KG] (approx. clove)" in the hidden `Recepies old`).
INGREDIENT_RENAMES = {
    "Basil ehm topf": "Basil (pot)",
    "Garlic ehm clove": "Garlic",
    "Chives ehm Topf": "Chives (pot)",
}
MARKER = re.compile(r"^!.*?!\s*")  # "!weg!", "!WEG! ", "!to be replaced! " mean: obsolete
IMPLAUSIBLE_RATIO = 3.0  # finished dish weighs this many times less than its ingredients -> needs review


def slug(text):
    text = text.replace("ß", "ss")
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return text


def cap(text):
    text = re.sub(r"\s+", " ", text).strip()
    return text[:1].upper() + text[1:]


def num(x, places):
    """Decimal -> plain string without trailing zeros or exponent."""
    q = Decimal(1).scaleb(-places)
    d = Decimal(x).quantize(q, rounding=ROUND_HALF_UP).normalize()
    s = format(d, "f")
    return s if s != "-0" else "0"


SAFE = re.compile(r"^[A-Za-z0-9ÄÖÜäöüß][A-Za-z0-9ÄÖÜäöüß .,()/&+'–-]*$")
RESERVED = {"yes", "no", "on", "off", "true", "false", "null", "y", "n", "~"}


def ys(text):
    """A YAML string: plain when that is safe, double-quoted otherwise."""
    if SAFE.match(text) and text.lower() not in RESERVED and not text.endswith(":") and " #" not in text:
        return text
    return json.dumps(text, ensure_ascii=False)


def main(path):
    wf = openpyxl.load_workbook(path)
    wv = openpyxl.load_workbook(path, data_only=True)
    rf, rv = wf["Recepies"], wv["Recepies"]
    table = rf.tables["Recepies"]
    ncol = len(table.tableColumns)
    nrow = int(re.search(r"(\d+)$", table.ref.split(":")[1]).group(1))
    report = collections.defaultdict(list)

    # --- ingredients: one per column between "Active Recepie" and "Working Steps" ---------------
    prices = {}
    ia = wv["Ingredient ALL"]
    for r in range(2, ia.max_row + 1):
        key, price = ia.cell(r, 1).value, ia.cell(r, 6).value
        if isinstance(key, str) and isinstance(price, (int, float)) and price > 0:
            prices[key] = price

    used = collections.Counter()
    for r in range(2, nrow + 1):
        for c in range(5, ncol):
            if rf.cell(r, c).value is not None:
                used[c] += 1

    ingredients = {}  # id -> dict
    col_to_id = {}
    for c in range(5, ncol):
        header = rf.cell(1, c).value
        raw_name, raw_unit, raw_store = [p.strip() for p in header.split("|")]
        if MARKER.match(raw_name):
            if not used[c]:
                report["Skipped ingredients (marked obsolete in the Excel and unused)"].append(header)
                continue
            report["Obsolete marker kept because a recipe uses it"].append(header)
        name = cap(INGREDIENT_RENAMES.get(MARKER.sub("", raw_name), MARKER.sub("", raw_name)))
        if raw_name in INGREDIENT_RENAMES:
            report["Renamed ingredients"].append(f"{raw_name} -> {name}")
        unit, mult = UNITS[raw_unit]
        store_id = STORES[raw_store][0]
        if raw_store == "SM2":
            report["Unknown store code treated as supermarket"].append(header)
        iid = slug(re.sub(r"\(.*?\)", "", name))  # "Red lentils (dry)" -> red-lentils
        if iid in ingredients:  # two names differ only in the brackets: keep the brackets in the id
            iid = slug(name)
        assert iid not in ingredients, f"duplicate ingredient id {iid}"
        entry = {"name": name, "unit": unit, "store": store_id}
        if header in prices and store_id != "none":
            entry["price"] = prices[header]
        else:
            if store_id != "none":
                report["Ingredients without a price"].append(name)
        ingredients[iid] = entry
        col_to_id[c] = (iid, mult, unit)
        if not used[c]:
            report["Ingredients no recipe uses (kept)"].append(name)

    # --- recipes ---------------------------------------------------------------------------------
    pat = re.compile(r"^=\s*(-?[0-9.]+(?:[eE][-+]?\d+)?)\s*\*")
    recipes = {}
    for r in range(2, nrow + 1):
        name = rv.cell(r, 1).value
        if not name or not str(name).strip():
            continue
        name = cap(str(name))
        rtype = (rv.cell(r, 2).value or "").strip().lower()
        grams = rv.cell(r, 3).value or 0
        lines = []
        for c in range(5, ncol):
            f = rf.cell(r, c).value
            if f is None or c not in col_to_id:
                continue
            m = pat.match(f) if isinstance(f, str) else None
            factor = Decimal(m.group(1)) if m else Decimal(str(f))
            iid, mult, unit = col_to_id[c]
            lines.append((iid, factor * mult, unit))
        if not lines:
            report["Skipped recipes (no ingredients)"].append(f"{name} (row {r})")
            continue
        rid = slug(re.sub(r"\(.*?\)", "", name))
        if rid in recipes:
            raise SystemExit(f"duplicate recipe id {rid}")
        weight = sum(a for _, a, u in lines if u in ("g", "ml"))  # g per 100 g of dish
        note = None
        if weight > IMPLAUSIBLE_RATIO * 100:
            note = (f"Ingredients add up to {num(weight, 0)} g per 100 g of finished dish "
                    f"({num(weight / 100, 1)} times). Probably entered per batch or in the wrong unit. Check amounts and units.")
        if name.lower().startswith("test"):
            note = (note + " " if note else "") + 'Called a test recipe in the Excel.'
        method = rv.cell(r, ncol).value
        method = str(method).strip() if method else ""
        if method.lower() == "kein rezept" or method.startswith("That is a pretty good question"):
            method = ""  # "no recipe", or the sheet's canned placeholder text
        lines.sort(key=lambda l: (-l[1] if l[2] in ("g", "ml") else Decimal(0), l[0]))
        recipes[rid] = {"name": name, "category": CATEGORIES[rtype], "portion": grams, "lines": lines,
                        "method": method, "note": note, "row": r}
        if note:
            report["Recipes marked needs-review"].append(f"{name}: {note}")
        if not method:
            report["Recipes without a method"].append(name)

    # --- write -----------------------------------------------------------------------------------
    DATA.mkdir(exist_ok=True)
    (DATA / "recipes").mkdir(exist_ok=True)
    for old in (DATA / "recipes").glob("*.yaml"):
        old.unlink()

    stores = {}
    for code, (sid, sname) in STORES.items():
        if sname:
            stores[sid] = sname
    out = ["# Where ingredients are bought. Every ingredient names exactly one store.",
           "# The store `none` is for things that are never bought, like tap water.", ""]
    out += [f"{ys(k)}: {ys(v)}" for k, v in stores.items()]
    (DATA / "stores.yaml").write_text("\n".join(out) + "\n")

    out = ["# Every ingredient the recipes can use.",
           "# unit:  how amounts are counted in recipes: g, ml, pcs, pack or jar.",
           "# price: euros per kilogram (unit g), per litre (unit ml) or per piece (pcs, pack, jar). Optional.",
           "# pack:  the usual pack size, like 25 kg or 400 ml. Optional.", ""]
    for iid in sorted(ingredients):
        e = ingredients[iid]
        out.append(f"{ys(iid)}:")
        out.append(f"  name: {ys(e['name'])}")
        out.append(f"  unit: {e['unit']}")
        out.append(f"  store: {e['store']}")
        if "price" in e:
            out.append(f"  price: {num(Decimal(str(e['price'])), 4)}")
    (DATA / "ingredients.yaml").write_text("\n".join(out) + "\n")

    for rid, rc in sorted(recipes.items()):
        out = [f"name: {ys(rc['name'])}", f"category: {rc['category']}"]
        if rc["portion"]:
            out.append(f"portion: {num(Decimal(str(rc['portion'])), 1)} g")
        out.append("source: Old Excel sheet")
        if rc["note"]:
            out.append("status: needs-review")
            out.append(f"review_note: {json.dumps(rc['note'], ensure_ascii=False)}")
        out.append("per_100g:")
        for iid, amount, unit in rc["lines"]:
            places = 4 if unit in ("g", "ml") else 6
            out.append(f"  {ys(iid)}: {num(amount, places)} {unit}")
        if rc["method"]:
            text = "\n".join(line.rstrip().replace("\t", "    ") for line in rc["method"].replace("\r", "").split("\n"))
            out.append("method: |")
            out += ["  " + line if line else "" for line in text.split("\n")]
        (DATA / "recipes" / f"{rid}.yaml").write_text("\n".join(out) + "\n")

    # --- report ----------------------------------------------------------------------------------
    lines = ["# Excel import report", "",
             "Generated by `import_excel.py`. Every guess the importer made is listed here.", "",
             f"- Recipes written: {len(recipes)}, of which marked needs-review: {len(report['Recipes marked needs-review'])}",
             f"- Ingredients written: {len(ingredients)}",
             f"- Stores: {len(stores)}", "",
             "Amounts: the Excel's kg and litres were converted to grams and millilitres (x1000). Pieces, packs and jars are unchanged.",
             "No ingredient has a pack size, because the Excel has none.", ""]
    for title in ["Recipes marked needs-review", "Skipped recipes (no ingredients)", "Skipped ingredients (marked obsolete in the Excel and unused)",
                  "Obsolete marker kept because a recipe uses it", "Renamed ingredients", "Unknown store code treated as supermarket",
                  "Recipes without a method", "Ingredients no recipe uses (kept)", "Ingredients without a price"]:
        items = report.get(title, [])
        lines += [f"## {title} ({len(items)})", ""] + [f"- {i}" for i in items] + [""]
    (HERE / "REPORT.md").write_text("\n".join(lines))
    print(f"recipes {len(recipes)}  ingredients {len(ingredients)}  needs-review {len(report['Recipes marked needs-review'])}")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit(__doc__)
    main(sys.argv[1])
