#!/usr/bin/env python3
"""Checks that the YAML written by import_excel.py says the same as the Excel.

Usage:  python verify_import.py PATH/TO/prev.xlsm

For every ingredient line of every imported recipe it recomputes the amount per person from the YAML
and compares it with the number Excel itself calculated in that cell. Needs: openpyxl, pyyaml.
"""
import re
import sys
import pathlib
import openpyxl
import yaml

ROOT = pathlib.Path(__file__).resolve().parent.parent.parent
DATA = ROOT / "data"
RENAMED = {"garlic": "garlic ehm clove", "basil (pot)": "basil ehm topf", "chives (pot)": "chives ehm topf"}  # see import_excel.py
TO_EXCEL_UNIT = {"g": 1000, "ml": 1000, "pcs": 1, "pack": 1, "jar": 1}  # YAML amount / this = Excel's kg, L or count


def qty(text):
    m = re.fullmatch(r"([0-9.]+) (g|kg|ml|l|pcs|pack|jar)", str(text))
    assert m, f"cannot read quantity {text!r}"
    return float(m.group(1)), m.group(2)


def main(path):
    ingredients = yaml.safe_load((DATA / "ingredients.yaml").read_text())
    stores = yaml.safe_load((DATA / "stores.yaml").read_text())
    ws_v = openpyxl.load_workbook(path, data_only=True)["Recepies"]
    headers = {}
    for c in range(5, ws_v.max_column):
        h = ws_v.cell(1, c).value
        if isinstance(h, str) and "|" in h:
            headers[re.sub(r"\s+", " ", h.split("|")[0].strip()).lower()] = c
    rows = {}
    for r in range(2, 149):
        n = ws_v.cell(r, 1).value
        if n and str(n).strip():
            rows.setdefault(re.sub(r"\s+", " ", str(n).strip()).lower(), []).append(r)

    problems = checked = 0
    for f in sorted((DATA / "recipes").glob("*.yaml")):
        rec = yaml.safe_load(f.read_text())
        assert rec["category"] in {"main", "brunch", "helper-snack", "drink", "misc"}, f.name
        portion = float(str(rec["portion"]).split()[0])
        r = rows[rec["name"].lower()][0]
        for iid, text in rec["per_100g"].items():
            ing = ingredients[iid]  # fails loudly if a recipe uses an unknown ingredient
            assert ing["store"] in stores, (iid, ing["store"])
            amount, unit = qty(text)
            assert unit == ing["unit"], f"{f.name}: {iid} is {ing['unit']} but written in {unit}"
            per_person_excel_unit = amount / TO_EXCEL_UNIT[unit] * portion / 100
            key = re.sub(r"\s+", " ", ing["name"]).lower()
            col = headers[RENAMED.get(key, key)]
            expected = ws_v.cell(r, col).value
            checked += 1
            if abs(per_person_excel_unit - float(expected)) > 1e-6 + 0.002 * abs(float(expected)):
                problems += 1
                print(f"MISMATCH {f.name} {iid}: yaml {per_person_excel_unit} vs excel {expected}")
    print(f"checked {checked} ingredient cells, {problems} mismatches")
    return problems


if __name__ == "__main__":
    raise SystemExit(1 if main(sys.argv[1]) else 0)
